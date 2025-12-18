"""Export service for generating documents from Q&A sessions."""

import logging
from typing import Dict, List, Any, Tuple
from datetime import datetime
from io import BytesIO

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.services.qa_storage import QAStorageService
from src.settings import Settings

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting Q&A sessions to various formats."""

    def __init__(self, settings: Settings):
        """
        Initialize export service.

        Args:
            settings: Application settings
        """
        self.settings = settings
        self.qa_storage: QAStorageService = None

    async def _get_qa_storage(self) -> QAStorageService:
        """Get or create QAStorageService instance."""
        if not self.qa_storage:
            self.qa_storage = QAStorageService(self.settings)
            await self.qa_storage.initialize()
        return self.qa_storage

    async def export_session_to_markdown(self, session_id: str) -> str:
        """
        Export session to Markdown format.

        Args:
            session_id: Session ID

        Returns:
            Markdown string

        Raises:
            ValueError: If session not found or has no Q&A pairs
        """
        qa_storage = await self._get_qa_storage()
        session = await qa_storage.get_session(session_id)
        qa_pairs = await qa_storage.get_session_qa_pairs(session_id)

        if not qa_pairs:
            raise ValueError(f"Session {session_id} has no Q&A pairs to export")

        # Build Markdown document
        lines = []
        
        # Header
        lines.append(f"# {session['session_name']}\n")
        
        # Session metadata
        metadata = session.get("metadata", {})
        company_info = metadata.get("company_info", {})
        
        if company_info:
            lines.append("## Company Information\n")
            if company_info.get("company_name"):
                lines.append(f"**Company:** {company_info['company_name']}\n")
            if company_info.get("industry"):
                lines.append(f"**Industry:** {company_info['industry']}\n")
            if company_info.get("activities"):
                lines.append(f"**Activities:** {company_info['activities']}\n")
            if company_info.get("context"):
                lines.append(f"**Context:** {company_info['context']}\n")
            lines.append("")
        
        # Session info
        lines.append("## Session Information\n")
        lines.append(f"**Created:** {session['created_at'].strftime('%Y-%m-%d %H:%M:%S UTC')}\n")
        lines.append(f"**Updated:** {session['updated_at'].strftime('%Y-%m-%d %H:%M:%S UTC')}\n")
        lines.append(f"**User Role:** {session.get('user_role', 'unknown').title()}\n")
        
        outcome_status = session.get("outcome_status")
        if outcome_status:
            lines.append(f"**Outcome Status:** {outcome_status.title()}\n")
        
        round_number = metadata.get("round_number", 1)
        if round_number > 1:
            lines.append(f"**Round:** {round_number}\n")
            parent_session_id = metadata.get("parent_session_id")
            if parent_session_id:
                lines.append(f"**Parent Session:** {parent_session_id}\n")
        
        lines.append("")
        
        # Q&A pairs
        lines.append("## Questions & Answers\n")
        lines.append("")
        
        for idx, qa_pair in enumerate(qa_pairs, 1):
            lines.append(f"### Question {idx}\n")
            lines.append(f"{qa_pair['question']}\n")
            lines.append("")
            lines.append("**Answer:**\n")
            lines.append(f"{qa_pair.get('final_answer', qa_pair.get('original_answer', ''))}\n")
            lines.append("")
            
            # Citations
            citations = qa_pair.get("citations", [])
            if citations:
                lines.append("**Citations:**\n")
                for cit_idx, citation in enumerate(citations, 1):
                    title = citation.get("title", citation.get("document_title", "Unknown"))
                    source = citation.get("source", "")
                    document_id = citation.get("document_id", "")
                    
                    # Format citation as numbered list item with link if source available
                    if source:
                        lines.append(f"{cit_idx}. [{title}]({source})")
                    else:
                        lines.append(f"{cit_idx}. {title}")
                    
                    if document_id:
                        lines.append(f"   (Document ID: {document_id})")
                    lines.append("")
            
            lines.append("---\n")
            lines.append("")
        
        return "\n".join(lines)

    async def export_session_to_pdf(self, session_id: str) -> bytes:
        """
        Export session to PDF format.

        Args:
            session_id: Session ID

        Returns:
            PDF bytes

        Raises:
            ValueError: If session not found or has no Q&A pairs
        """
        qa_storage = await self._get_qa_storage()
        session = await qa_storage.get_session(session_id)
        qa_pairs = await qa_storage.get_session_qa_pairs(session_id)

        if not qa_pairs:
            raise ValueError(f"Session {session_id} has no Q&A pairs to export")

        # Create PDF buffer
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor='#000000',
            spaceAfter=12,
            alignment=TA_LEFT
        )
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor='#333333',
            spaceAfter=8,
            spaceBefore=12,
            alignment=TA_LEFT
        )
        question_style = ParagraphStyle(
            'QuestionStyle',
            parent=styles['Normal'],
            fontSize=12,
            textColor='#000000',
            fontName='Helvetica-Bold',
            spaceAfter=6,
            alignment=TA_LEFT
        )
        answer_style = ParagraphStyle(
            'AnswerStyle',
            parent=styles['Normal'],
            fontSize=11,
            textColor='#333333',
            spaceAfter=12,
            alignment=TA_LEFT
        )
        citation_style = ParagraphStyle(
            'CitationStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor='#666666',
            leftIndent=0.5*inch,
            spaceAfter=4,
            alignment=TA_LEFT
        )
        
        # Build PDF content
        story = []
        
        # Title
        story.append(Paragraph(session['session_name'], title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Session metadata
        metadata = session.get("metadata", {})
        company_info = metadata.get("company_info", {})
        
        if company_info:
            story.append(Paragraph("Company Information", heading_style))
            if company_info.get("company_name"):
                story.append(Paragraph(f"<b>Company:</b> {company_info['company_name']}", styles['Normal']))
            if company_info.get("industry"):
                story.append(Paragraph(f"<b>Industry:</b> {company_info['industry']}", styles['Normal']))
            if company_info.get("activities"):
                story.append(Paragraph(f"<b>Activities:</b> {company_info['activities']}", styles['Normal']))
            if company_info.get("context"):
                story.append(Paragraph(f"<b>Context:</b> {company_info['context']}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
        
        # Session info
        story.append(Paragraph("Session Information", heading_style))
        story.append(Paragraph(f"<b>Created:</b> {session['created_at'].strftime('%Y-%m-%d %H:%M:%S UTC')}", styles['Normal']))
        story.append(Paragraph(f"<b>Updated:</b> {session['updated_at'].strftime('%Y-%m-%d %H:%M:%S UTC')}", styles['Normal']))
        story.append(Paragraph(f"<b>User Role:</b> {session.get('user_role', 'unknown').title()}", styles['Normal']))
        
        outcome_status = session.get("outcome_status")
        if outcome_status:
            story.append(Paragraph(f"<b>Outcome Status:</b> {outcome_status.title()}", styles['Normal']))
        
        round_number = metadata.get("round_number", 1)
        if round_number > 1:
            story.append(Paragraph(f"<b>Round:</b> {round_number}", styles['Normal']))
        
        story.append(Spacer(1, 0.2*inch))
        
        # Q&A pairs
        story.append(Paragraph("Questions & Answers", heading_style))
        story.append(Spacer(1, 0.1*inch))
        
        for idx, qa_pair in enumerate(qa_pairs, 1):
            # Question
            question_text = qa_pair['question'].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(f"Question {idx}", heading_style))
            story.append(Paragraph(question_text, question_style))
            
            # Answer
            answer = qa_pair.get('final_answer', qa_pair.get('original_answer', ''))
            answer_text = answer.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph("<b>Answer:</b>", styles['Normal']))
            story.append(Paragraph(answer_text, answer_style))
            
            # Citations
            citations = qa_pair.get("citations", [])
            if citations:
                story.append(Paragraph("<b>Citations:</b>", styles['Normal']))
                for cit_idx, citation in enumerate(citations, 1):
                    title = citation.get("title", citation.get("document_title", "Unknown"))
                    source = citation.get("source", "")
                    document_id = citation.get("document_id", "")
                    
                    cit_text = f"{cit_idx}. {title}"
                    if source:
                        cit_text += f" ({source})"
                    if document_id:
                        cit_text += f" [ID: {document_id}]"
                    
                    cit_text = cit_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(cit_text, citation_style))
            
            story.append(Spacer(1, 0.2*inch))
        
        # Build PDF
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    async def export_session_to_word(self, session_id: str) -> bytes:
        """
        Export session to Word format.

        Args:
            session_id: Session ID

        Returns:
            DOCX bytes

        Raises:
            ValueError: If session not found or has no Q&A pairs
        """
        qa_storage = await self._get_qa_storage()
        session = await qa_storage.get_session(session_id)
        qa_pairs = await qa_storage.get_session_qa_pairs(session_id)

        if not qa_pairs:
            raise ValueError(f"Session {session_id} has no Q&A pairs to export")

        # Create Word document
        doc = Document()
        
        # Title
        title = doc.add_heading(session['session_name'], 0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Session metadata
        metadata = session.get("metadata", {})
        company_info = metadata.get("company_info", {})
        
        if company_info:
            doc.add_heading('Company Information', 1)
            if company_info.get("company_name"):
                p = doc.add_paragraph()
                p.add_run('Company: ').bold = True
                p.add_run(company_info['company_name'])
            if company_info.get("industry"):
                p = doc.add_paragraph()
                p.add_run('Industry: ').bold = True
                p.add_run(company_info['industry'])
            if company_info.get("activities"):
                p = doc.add_paragraph()
                p.add_run('Activities: ').bold = True
                p.add_run(company_info['activities'])
            if company_info.get("context"):
                p = doc.add_paragraph()
                p.add_run('Context: ').bold = True
                p.add_run(company_info['context'])
        
        # Session info
        doc.add_heading('Session Information', 1)
        p = doc.add_paragraph()
        p.add_run('Created: ').bold = True
        p.add_run(session['created_at'].strftime('%Y-%m-%d %H:%M:%S UTC'))
        
        p = doc.add_paragraph()
        p.add_run('Updated: ').bold = True
        p.add_run(session['updated_at'].strftime('%Y-%m-%d %H:%M:%S UTC'))
        
        p = doc.add_paragraph()
        p.add_run('User Role: ').bold = True
        p.add_run(session.get('user_role', 'unknown').title())
        
        outcome_status = session.get("outcome_status")
        if outcome_status:
            p = doc.add_paragraph()
            p.add_run('Outcome Status: ').bold = True
            p.add_run(outcome_status.title())
        
        round_number = metadata.get("round_number", 1)
        if round_number > 1:
            p = doc.add_paragraph()
            p.add_run('Round: ').bold = True
            p.add_run(str(round_number))
        
        # Q&A pairs
        doc.add_heading('Questions & Answers', 1)
        
        for idx, qa_pair in enumerate(qa_pairs, 1):
            # Question
            doc.add_heading(f'Question {idx}', 2)
            doc.add_paragraph(qa_pair['question'])
            
            # Answer
            p = doc.add_paragraph()
            p.add_run('Answer: ').bold = True
            answer = qa_pair.get('final_answer', qa_pair.get('original_answer', ''))
            doc.add_paragraph(answer)
            
            # Citations
            citations = qa_pair.get("citations", [])
            if citations:
                p = doc.add_paragraph()
                p.add_run('Citations:').bold = True
                for cit_idx, citation in enumerate(citations, 1):
                    p = doc.add_paragraph(f"{cit_idx}. ", style='List Number')
                    title = citation.get("title", citation.get("document_title", "Unknown"))
                    p.add_run(title)
                    source = citation.get("source", "")
                    if source:
                        p.add_run(f" ({source})")
                    document_id = citation.get("document_id", "")
                    if document_id:
                        p.add_run(f" [ID: {document_id}]")
        
        # Save to bytes buffer
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    async def export_session(
        self, session_id: str, format: str
    ) -> Tuple[bytes, str, str]:
        """
        Export session to specified format.

        Args:
            session_id: Session ID
            format: Export format ("markdown", "pdf", or "docx")

        Returns:
            Tuple of (file_bytes, mime_type, filename)

        Raises:
            ValueError: If format is invalid or session not found
        """
        format_lower = format.lower()
        
        if format_lower == "markdown":
            content = await self.export_session_to_markdown(session_id)
            file_bytes = content.encode('utf-8')
            mime_type = "text/markdown"
            extension = "md"
        elif format_lower == "pdf":
            file_bytes = await self.export_session_to_pdf(session_id)
            mime_type = "application/pdf"
            extension = "pdf"
        elif format_lower == "docx":
            file_bytes = await self.export_session_to_word(session_id)
            mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            extension = "docx"
        else:
            raise ValueError(f"Unsupported export format: {format}")

        # Generate filename
        qa_storage = await self._get_qa_storage()
        session = await qa_storage.get_session(session_id)
        session_name = session['session_name']
        # Sanitize filename
        safe_name = "".join(c for c in session_name if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_name = safe_name.replace(' ', '_')
        timestamp = datetime.utcnow().strftime('%Y%m%d_%H%M%S')
        filename = f"{safe_name}_{timestamp}.{extension}"

        return file_bytes, mime_type, filename

